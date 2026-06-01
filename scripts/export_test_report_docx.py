from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "TEST_CASE_TREE.md"
OUTPUT = ROOT / "TEST_CASE_TREE.docx"


def clean_text(value: str) -> str:
    text = value.replace("<br>", "\n")
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = text.replace("\\|", "|")
    return text.strip()


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        if index < len(lines) - 1:
            run.add_break()


def add_table(document: Document, rows: list[str]) -> None:
    parsed_rows = []
    for row in rows:
        if set(row.replace("|", "").replace("-", "").replace(" ", "")) == set():
            continue
        cells = [clean_text(cell) for cell in row.strip().strip("|").split("|")]
        parsed_rows.append(cells)

    if not parsed_rows:
        return

    table = document.add_table(rows=len(parsed_rows), cols=len(parsed_rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for row_index, row in enumerate(parsed_rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    if row_index == 0:
                        run.bold = True

    document.add_paragraph("")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)


def main() -> None:
    document = Document()
    configure_document(document)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0

    while index < len(lines):
        line = lines[index]

        if not line.strip():
            document.add_paragraph("")
            index += 1
            continue

        if line.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        if line.startswith("|"):
            table_lines = [line]
            index += 1
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = clean_text(line[level:].strip())
            document.add_heading(text, level=min(level, 4))
            index += 1
            continue

        if line.startswith("- "):
            document.add_paragraph(clean_text(line[2:]), style="List Bullet")
            index += 1
            continue

        if re.match(r"^\d+\.\s", line):
            document.add_paragraph(clean_text(re.sub(r"^\d+\.\s*", "", line)), style="List Number")
            index += 1
            continue

        document.add_paragraph(clean_text(line))
        index += 1

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
